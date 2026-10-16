"""Load and clean spend workbooks. Writes outputs/analytics-full/{cleaned_spend,excluded_rows,public_body_matches,public_body_review}.csv
and outputs/analytics-full/_clean_meta.json."""
import glob, json, os, re, sys
from datetime import datetime
import openpyxl
import pandas as pd

# ---------------- config ----------------
EXPECTED = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]
SHEET = "Data to publish"
PENSION_PATTERNS = [r"LGPS", r"\bAVC\b", r"Pension", r"Court Order", r"Council Tax"]
PUBLIC_BODY_PATTERNS = [r"\bCOUNCIL\b", r"\bCNCL\b", r"^ROYAL BOROUGH OF", r"\bNHS\b", r"\bICB\b",
                        r"INTEGRATED CARE BOARD", r"\bPOLICE\b", r"FIRE (?:AND|&) RESCUE", r"\bHMRC\b",
                        r"REVENUE (?:AND|&) CUSTOMS", r"^MINISTRY OF", r"^DEPARTMENT (?:FOR|OF)",
                        r"COMBINED AUTHORITY", r"HOMES ENGLAND"]
PUBLIC_BODY_NARRATIVE = r"^Joint Arrangements"
KEEP_AS_SUPPLIER = []  # filled after reviewing public_body_matches.csv
REVIEW_PATTERNS = [r"\bSCHOOLS?\b", r"\bACADEM(?:Y|IES)\b", r"\bCOLLEGE\b", r"\bUNIVERSITY\b", r"\bTRUST\b"]
TAG_AGENCY = r"agency\s*&?\s*temp"
TAG_GRANT = r"\bgrants?\b"
TAG_PREMISES_CAT = "Premises"
TAG_PREMISES_NARR = r"R&M|minor works|grounds maintenance|cleaning|\brents?\b|repairs"
CARE_SERVICES = ["Adult Social Care", "Children's Social Care", "Education & SEND",
                 "Education (DSG Funded) & Other Education Grants"]
CARE_CATEGORY = "Third Party Payments"
OUT = "outputs/analytics-full"


def norm(s):
    s = str(s).upper()
    s = re.sub(r"[^\w\s]|_", " ", s)
    s = re.sub(r"\b(LTD|LIMITED|PLC|LLP)\b", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def load():
    frames, files = [], []
    for f in sorted(glob.glob("data/spend/*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            sys.exit(f"STOP: {f} lacks sheet '{SHEET}'")
        ws = wb[SHEET]
        rows = list(ws.iter_rows(values_only=True))
        hdr = None
        for i, r in enumerate(rows[:15]):
            cells = [str(c).strip() if c is not None else "" for c in r]
            if all(c in cells for c in EXPECTED):
                hdr = i; idx = [cells.index(c) for c in EXPECTED]; break
        if hdr is None:
            seen = set(str(c).strip() for r in rows[:15] for c in r if c is not None)
            sys.exit(f"STOP: {f}: header not found in first 15 rows; missing {[c for c in EXPECTED if c not in seen]}")
        stem = os.path.splitext(os.path.basename(f))[0]
        recs, blank = [], 0
        for j in range(hdr + 1, len(rows)):
            vals = [rows[j][k] if k < len(rows[j]) else None for k in idx]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
                blank += 1; continue
            recs.append([f"{stem}|{SHEET}|{j + 1}", stem] + vals)
        files.append({"file": os.path.basename(f), "sheet": SHEET, "header_row": hdr + 1,
                      "rows_read": len(recs), "rows_blank": blank})
        frames.append(pd.DataFrame(recs, columns=["row_id", "file", "service", "category", "narrative", "date", "amount", "supplier"]))
    df = pd.concat(frames, ignore_index=True)
    # dates
    def pd_date(v):
        if isinstance(v, datetime): return v
        if isinstance(v, str):
            try: return datetime.strptime(v.strip(), "%m/%d/%y")
            except ValueError: sys.exit(f"STOP: unparseable date {v!r}")
        sys.exit(f"STOP: bad date value {v!r}")
    df["date"] = pd.to_datetime(df["date"].map(pd_date))
    amt = pd.to_numeric(df["amount"], errors="coerce")
    if amt.isna().any():
        sys.exit(f"STOP: {int(amt.isna().sum())} non-numeric Net amount values")
    df["amount"] = amt.astype(float)
    for c in ["service", "category", "narrative", "supplier"]:
        df[c] = df[c].fillna("").astype(str).str.strip()
    df["norm_supplier"] = df["supplier"].map(norm)
    return df, files


def exclusion(df):
    reason = pd.Series("", index=df.index)
    def setr(mask, name):
        reason[(reason == "") & mask] = name
    setr(df.supplier.str.lower() == "redacted", "redacted_supplier")
    setr(df.narrative.str.contains("|".join(PENSION_PATTERNS), case=False, regex=True), "pension_statutory")
    setr((df.category == "Transfer Payment") | (df.narrative == "ASC Deferred Debt"), "personal_transfer")
    pb = df.supplier.str.upper().str.contains("|".join(PUBLIC_BODY_PATTERNS), regex=True) | \
         df.narrative.str.contains(PUBLIC_BODY_NARRATIVE, case=False, regex=True)
    keep = df.supplier.str.upper().isin([k.upper() for k in KEEP_AS_SUPPLIER])
    df["_pb_match"] = pb
    setr(pb & ~keep, "public_body")
    return reason


def tag(r):
    if re.search(TAG_AGENCY, r.narrative, re.I): return "agency_staff"
    if re.search(TAG_GRANT, r.narrative, re.I): return "grant"
    if r.category == TAG_PREMISES_CAT or re.search(TAG_PREMISES_NARR, r.narrative, re.I): return "premises"
    if r.service in CARE_SERVICES and r.category == CARE_CATEGORY: return "care_education_placement"
    return "general_procurement"


def main():
    os.makedirs(OUT, exist_ok=True)
    df, files = load()
    df["exclusion_reason"] = exclusion(df)
    ex = df[df.exclusion_reason != ""]
    ex[["row_id", "exclusion_reason", "supplier", "narrative", "category", "amount"]].rename(
        columns={"exclusion_reason": "reason"}).to_csv(f"{OUT}/excluded_rows.csv", index=False)
    pbm = df[df._pb_match].groupby("supplier").agg(gbp=("amount", "sum"), rows=("amount", "size")).reset_index().sort_values("gbp", ascending=False)
    pbm["kept_as_supplier"] = pbm.supplier.str.upper().isin([k.upper() for k in KEEP_AS_SUPPLIER])
    pbm.to_csv(f"{OUT}/public_body_matches.csv", index=False)
    k = df[df.exclusion_reason == ""].copy()
    rv = k[k.supplier.str.upper().str.contains("|".join(REVIEW_PATTERNS), regex=True)]
    rv.groupby("supplier").amount.sum().round(2).reset_index().rename(columns={"amount": "gbp"}).sort_values("gbp", ascending=False).to_csv(f"{OUT}/public_body_review.csv", index=False)
    k["tag"] = k.apply(tag, axis=1)
    cols = ["row_id", "file", "date", "service", "category", "narrative", "amount", "supplier", "norm_supplier", "tag"]
    k[cols].to_csv(f"{OUT}/cleaned_spend.csv", index=False)
    months = k.date.dt.to_period("M").nunique()
    ws = k.date.min().replace(day=1).normalize()
    we = (k.date.max() + pd.offsets.MonthEnd(0)).normalize()
    meta = {"spend_files": files, "rows_read": int(len(df)),
            "excluded": {r: {"rows": int(g.shape[0]), "gbp": round(float(g.amount.sum()), 2)} for r, g in ex.groupby("exclusion_reason")},
            "rows_analysed": int(len(k)), "gbp_analysed": round(float(k.amount.sum()), 2),
            "suppliers_analysed": int(k.norm_supplier.nunique()),
            "tags": {t: {"rows": int(len(g)), "gbp": round(float(g.amount.sum()), 2), "suppliers": int(g.norm_supplier.nunique())} for t, g in k.groupby("tag")},
            "window": {"start": ws.date().isoformat(), "end": we.date().isoformat(), "months": int(months), "annualise_factor": 12 / months},
            "exclusion_patterns": {"pension_statutory": PENSION_PATTERNS, "public_body": PUBLIC_BODY_PATTERNS, "public_body_narrative": PUBLIC_BODY_NARRATIVE,
                                   "personal_transfer": ["category == Transfer Payment", "narrative == ASC Deferred Debt"]},
            "keep_as_supplier": KEEP_AS_SUPPLIER}
    json.dump(meta, open(f"{OUT}/_clean_meta.json", "w"), indent=1)
    print(json.dumps({k_: meta[k_] for k_ in ["rows_read", "excluded", "rows_analysed", "gbp_analysed", "suppliers_analysed", "tags", "window"]}, indent=1))
    print(pd.crosstab(k.service, k.tag).head(25).to_string())

if __name__ == "__main__":
    main()
