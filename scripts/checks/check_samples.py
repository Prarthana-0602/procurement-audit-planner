#!/usr/bin/env python3
"""Sample check. Every transaction in outputs/samples/*.csv, and every row id in
audit-program.md, must (1) exist in outputs/analytics-full/cleaned_spend.csv with the
same supplier and amount, and (2) exist at that sheet row of the council's workbook in
data/spend/ with the same supplier and amount. The sheet row is the last part of the row
id: <file stem>|Data to publish|<excel row>.
    .venv/bin/python scripts/checks/check_samples.py [--root DIR]
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import ROW_ID_RE, cli, read_text, result  # noqa: E402

SHEET = "Data to publish"
EXPECTED = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]


def money(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def workbook_rows(path, wanted, items):
    """{excel row: (supplier, amount)} for the wanted rows. Counts rows like scripts/clean.py."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        items.append(f"{path.name}: no sheet {SHEET!r}")
        return {}
    rows = list(wb[SHEET].iter_rows(values_only=True))
    for r in rows[:15]:
        cells = [str(c).strip() if c is not None else "" for c in r]
        if all(c in cells for c in EXPECTED):
            si, ai = cells.index("Supplier name"), cells.index("Net amount")
            break
    else:
        items.append(f"{path.name}: header row with {EXPECTED} not found in the first 15 rows")
        return {}
    out = {}
    for n in wanted:
        if 1 <= n <= len(rows):
            r = rows[n - 1]
            out[n] = (str(r[si]).strip() if si < len(r) and r[si] is not None else "", money(r[ai]) if ai < len(r) else None)
    return out


def check(root):
    items = []
    out = root / "outputs"
    spend_path = out / "analytics-full" / "cleaned_spend.csv"
    if not spend_path.exists():
        return result("check_samples", [f"cleaned_spend.csv not found: {spend_path}"], "")
    with open(spend_path, newline="", encoding="utf-8") as fh:
        spend = {r["row_id"]: (r["supplier"].strip(), money(r["amount"])) for r in csv.DictReader(fh)}
    want = {}  # row_id -> where it was asked for
    files = sorted((out / "samples").glob("*.csv")) if (out / "samples").is_dir() else []
    if not files:
        items.append(f"no sample files in {out / 'samples'}")
    for f in files:
        with open(f, newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        if not rows:
            items.append(f"{f.name}: empty sample")
        for i, r in enumerate(rows, 2):
            rid = r.get("row_id", "")
            want.setdefault(rid, []).append(f"{f.name} line {i}")
            if rid in spend and ((r.get("supplier") or "").strip(), money(r.get("amount"))) != spend[rid]:
                items.append(f"{f.name} line {i}: {rid} supplier/amount differ from cleaned_spend.csv "
                             f"({r.get('supplier')!r}, {r.get('amount')} vs {spend[rid][0]!r}, {spend[rid][1]})")
    program = read_text(out / "audit-program.md", items, "audit-program.md")
    for m in ROW_ID_RE.finditer(program or ""):
        want.setdefault(m.group(0), []).append("audit-program.md")
    for rid, where in want.items():
        if rid not in spend:
            items.append(f"{rid} ({where[0]}): not in cleaned_spend.csv")
    by_file = {}
    for rid in want:
        m = re.fullmatch(r"(.+)\|" + re.escape(SHEET) + r"\|(\d+)", rid)
        if m:
            by_file.setdefault(m.group(1), set()).add(int(m.group(2)))
        else:
            items.append(f"{rid}: not a row id (<file stem>|{SHEET}|<excel row>)")
    for stem, wanted in sorted(by_file.items()):
        wbp = root / "data" / "spend" / f"{stem}.xlsx"
        if not wbp.exists():
            items.append(f"{stem}: workbook not found: {wbp}")
            continue
        got = workbook_rows(wbp, wanted, items)
        for n in sorted(wanted):
            rid = f"{stem}|{SHEET}|{n}"
            if n not in got:
                items.append(f"{rid} ({want[rid][0]}): no such row in the workbook")
            elif rid in spend and got[n] != spend[rid]:
                items.append(f"{rid} ({want[rid][0]}): workbook row has {got[n][0]!r}, {got[n][1]}; cleaned_spend.csv has "
                             f"{spend[rid][0]!r}, {spend[rid][1]}")
    return result("check_samples", items, f"{len(want)} distinct transaction(s) in {len(files)} sample file(s) and the program "
                  "exist in cleaned_spend.csv and at their sheet row in the workbooks, same supplier and amount")


if __name__ == "__main__":
    sys.exit(cli(check, __doc__.split("\n\n")[0]))
